
from pathlib import Path

from openpyxl import load_workbook

from src.combine_detail import extract_not_match_by_file_name
from src.excel_filter import (
    MASTER_FILE_STEM,
    display_group_title,
    get_unique_field_names_with_not_match,
)
from src.excel_exporter import export_round


class ExcelFilterLibrary:
    ROBOT_LIBRARY_SCOPE = "SUITE"

    def __init__(self):
        self.project_dir = Path(__file__).resolve().parents[1]
        self.last_output = None
        self.last_metadata = None
        self.last_round = None

    def _get_round_input_dir(self, round_name="Round14"):
        round_name = str(round_name).strip()

        if (
            not round_name
            or Path(round_name).name != round_name
            or "/" in round_name
            or "\\" in round_name
            or round_name in {".", ".."}
        ):
            raise ValueError(f"Invalid Round name: {round_name!r}")

        input_dir = self.project_dir / "input" / round_name

        if not input_dir.is_dir():
            raise AssertionError(
                f"Input folder not found: {input_dir}"
            )

        return input_dir

    def validate_round_input_folder(self, round_name="Round14"):
        return str(self._get_round_input_dir(round_name))

    def process_all_excel_files(self, round_name="Round14"):
        """
        Process Excel files only in input/<Round>.
        Output: output/<Round>_Combined_Filtered.xlsx
        """

        round_name = str(round_name).strip()
        self._get_round_input_dir(round_name)
        metadata = export_round(round_name)
        output_file = Path(metadata["output_file"])

        self.last_output = output_file
        self.last_metadata = metadata
        self.last_round = round_name

        return str(output_file)

    def get_last_export_summary(self):
        if self.last_metadata is None:
            raise AssertionError(
                "Run Process All Excel Files before requesting the export summary."
            )

        detail_counts = [
            {
                "name": display_group_title(record["path"]),
                "count": len(record["values"]),
            }
            for record in self.last_metadata["records"]
            if not record["is_master"]
        ]

        return {
            "round_name": self.last_round,
            "input_dir": self.last_metadata["input_dir"],
            "output_file": str(self.last_output),
            "master_name": display_group_title(
                self.last_metadata["master_file"]
            ),
            "master_count": self.last_metadata["master_count"],
            "detail_file_count": len(
                self.last_metadata["detail_files"]
            ),
            "detail_counts": detail_counts,
        }

    def get_last_field_name_counts(self):
        if self.last_metadata is None:
            raise AssertionError(
                "Run Process All Excel Files before requesting field name counts."
            )

        file_counts = []
        for record in self.last_metadata["records"]:
            file_path = Path(record["path"])
            field_names = get_unique_field_names_with_not_match(
                file_path
            )
            file_counts.append(
                {
                    "type": "Master" if record["is_master"] else "Detail",
                    "name": display_group_title(file_path),
                    "count": len(field_names),
                }
            )

        return file_counts

    def get_input_file_count(self, round_name="Round14"):
        return len(self._get_input_excel_files(round_name))

    def get_input_file_names(self, round_name="Round14"):
        return [
            path.name
            for path in self._get_input_excel_files(round_name)
        ]

    def get_master_detail_extraction_counts(self, round_name="Round14"):
        excel_files = self._get_input_excel_files(round_name)
        master_files = [
            path
            for path in excel_files
            if MASTER_FILE_STEM.casefold() in path.stem.casefold()
        ]

        if not master_files:
            raise FileNotFoundError(
                f"Master file containing '{MASTER_FILE_STEM}' "
                f"was not found for {round_name}."
            )

        if len(master_files) > 1:
            raise ValueError(
                f"Multiple Master files found for {round_name}: "
                f"{[path.name for path in master_files]}"
            )

        master_file = master_files[0]
        detail_files = [
            path for path in excel_files if path != master_file
        ]

        if not detail_files:
            raise FileNotFoundError(
                f"No Detail Excel files found for {round_name}."
            )

        file_counts = []
        for file_path in [master_file, *detail_files]:
            groups = [
                {
                    "file_name": group["file_name"],
                    "count": len(group["values"]),
                }
                for group in extract_not_match_by_file_name(file_path)
            ]
            file_counts.append(
                {
                    "type": (
                        "Master" if file_path == master_file else "Detail"
                    ),
                    "file_name": (
                        display_group_title(file_path)
                        if file_path == master_file
                        else file_path.name
                    ),
                    "count": sum(group["count"] for group in groups),
                    "groups": groups,
                }
            )

        return file_counts

    def _get_input_excel_files(self, round_name):
        input_dir = self._get_round_input_dir(round_name)

        return sorted(
            [
                path
                for path in input_dir.iterdir()
                if path.is_file()
                and path.suffix.lower() in {".xlsx", ".xlsm"}
                and not path.name.startswith("~$")
            ],
            key=lambda path: path.name.casefold(),
        )

    def verify_output_files_exist(self, round_name=None):
        if self.last_output is None:
            raise AssertionError(
                "Output has not been generated yet."
            )

        if round_name is not None:
            if str(round_name).strip() != self.last_round:
                raise AssertionError(
                    f"Selected Round is {round_name}, "
                    f"but last exported Round is {self.last_round}."
                )

        if not self.last_output.is_file():
            raise AssertionError(
                f"Output file does not exist: {self.last_output}"
            )

        if self.last_output.stat().st_size == 0:
            raise AssertionError(
                f"Output file is empty: {self.last_output}"
            )

        return True

    def validate_excel_comparison_results(self):
        if self.last_output is None:
            raise AssertionError(
                "Run Process All Excel Files first."
            )

        if not self.last_output.is_file():
            raise AssertionError(
                f"Output file not found: {self.last_output}"
            )

        workbook = load_workbook(
            self.last_output,
            read_only=False,
            data_only=False,
        )

        try:
            required_sheets = {
                "Summary",
                "Combine result",
            }

            missing_sheets = (
                required_sheets - set(workbook.sheetnames)
            )

            if missing_sheets:
                raise AssertionError(
                    f"Missing sheets: {missing_sheets}"
                )

            ws = workbook["Combine result"]
            status_count = 0

            for row in ws.iter_rows():
                for cell in row:
                    if cell.value not in ("Match", "Diff"):
                        continue

                    status_count += 1
                    color = cell.fill.fgColor.rgb

                    if cell.value == "Match":
                        valid_colors = {
                            "FFD9EAD3",
                            "D9EAD3",
                        }

                        if color not in valid_colors:
                            raise AssertionError(
                                f"Match cell {cell.coordinate} "
                                f"does not have green fill. "
                                f"Found: {color}"
                            )

                    elif cell.value == "Diff":
                        valid_colors = {
                            "FFF4CCCC",
                            "F4CCCC",
                        }

                        if color not in valid_colors:
                            raise AssertionError(
                                f"Diff cell {cell.coordinate} "
                                f"does not have red fill. "
                                f"Found: {color}"
                            )

            if status_count == 0:
                raise AssertionError(
                    "No Match/Diff status cells found."
                )

        finally:
            workbook.close()

        return True