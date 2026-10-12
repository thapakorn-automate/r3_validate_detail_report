import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.excel_filter import (
    build_combine_detail_sheet,
    build_comparison_sheet,
)


class CombineDetailTests(unittest.TestCase):
    def create_report(self, file_path, rows):
        workbook = Workbook()
        worksheet = workbook.active

        for row in rows:
            worksheet.append(row)

        workbook.save(file_path)
        workbook.close()

    def test_combines_unique_not_match_values_by_file_name(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            master_report = (
                root / "Detail_Normal_14_999_999_260922_Diff.xlsx"
            )
            first_report = (
                root / "Detail_Normal_14_470_991_260922_Diff.xlsx"
            )
            second_report = (
                root / "Detail_Normal_14_471_991_260922_Diff.xlsx"
            )

            self.create_report(
                master_report,
                [
                    ["File Name", None, "Rate Index Code Options"],
                    [None, None, None, None, None, "Result"],
                    [None, None, "VALUE_A", None, None, "Not match"],
                    [None, None, "MASTER_ONLY", None, None, "Not match"],
                    ["File Name", None, "GL Prod/Fis Code Lookup"],
                    [None, None, None, None, None, "Result"],
                    [None, None, "VALUE_X", None, None, "Not match"],
                ],
            )
            self.create_report(
                first_report,
                [
                    ["File Name", None, "Rate Index Code Options"],
                    [None, None, None, None, None, "Result"],
                    [None, None, "MATCH_VALUE", None, None, "Match"],
                    [None, None, "VALUE_A", None, None, "Not match"],
                    [None, None, "VALUE_A", None, None, "Not match"],
                    [None, None, "VALUE_B", None, None, "Not match"],
                    ["File Name", None, "GL Prod/Fis Code Lookup"],
                    [None, None, None, None, None, "Result"],
                    [None, None, "VALUE_X", None, None, "Not match"],
                ],
            )
            self.create_report(
                second_report,
                [
                    ["File Name", None, "Rate Index Code Options"],
                    [None, None, None, None, None, "Result"],
                    [None, None, "value_a", None, None, "Not match"],
                    [None, None, "VALUE_C", None, None, "Not match"],
                ],
            )

            output_file = root / "combined.xlsx"
            output_workbook = Workbook()
            build_combine_detail_sheet(
                output_workbook.active,
                [
                    {"path": second_report},
                    {"path": master_report},
                    {"path": first_report},
                ],
            )
            output_workbook.save(output_file)
            output_workbook.close()

            result_workbook = load_workbook(
                output_file,
                read_only=True,
                data_only=True,
            )
            try:
                worksheet = result_workbook.active
                first_detail_values = [
                    (
                        worksheet.cell(row, 4).value,
                        worksheet.cell(row, 5).value,
                    )
                    for row in range(3, worksheet.max_row + 1)
                    if worksheet.cell(row, 4).value is not None
                    and worksheet.cell(row, 5).value is not None
                    and worksheet.cell(row, 4).value != "Count"
                ]
                first_detail_statuses = [
                    (
                        worksheet.cell(row, 6).value,
                        worksheet.cell(row, 7).value,
                    )
                    for row in range(3, 7)
                ]

                self.assertEqual(
                    worksheet.cell(1, 1).value,
                    "14_999_999",
                )
                self.assertEqual(
                    worksheet.cell(2, 1).value,
                    "File Name",
                )
                self.assertEqual(
                    worksheet.cell(2, 2).value,
                    "Not match value",
                )
                header_cell = worksheet.cell(2, 1)
                group_title_cell = worksheet.cell(1, 1)
                self.assertEqual(
                    header_cell.fill.fgColor.rgb,
                    group_title_cell.fill.fgColor.rgb,
                )
                self.assertTrue(
                    header_cell.fill.fgColor.rgb.endswith("1F4E78")
                )
                self.assertEqual(header_cell.font.name, "Arial")
                self.assertEqual(header_cell.font.bold, True)
                self.assertEqual(header_cell.font.color.type, "rgb")
                self.assertTrue(
                    header_cell.font.color.rgb.endswith("FFFFFF")
                )
                self.assertEqual(
                    worksheet.cell(3, 4).font.name,
                    "Arial",
                )
                self.assertEqual(
                    worksheet.cell(3, 4).font.sz,
                    10,
                )
                self.assertEqual(
                    header_cell.alignment.horizontal,
                    "center",
                )
                self.assertEqual(
                    worksheet.cell(3, 4).alignment.horizontal,
                    "left",
                )
                self.assertTrue(
                    worksheet.cell(3, 6).fill.fgColor.rgb.endswith(
                        "D9EAD3"
                    )
                )
                self.assertTrue(
                    worksheet.cell(4, 6).fill.fgColor.rgb.endswith(
                        "F4CCCC"
                    )
                )
                self.assertEqual(
                    first_detail_values,
                    [
                        ("Rate Index Code Options", "VALUE_A"),
                        ("Rate Index Code Options", "VALUE_B"),
                        ("GL Prod/Fis Code Lookup", "VALUE_X"),
                        ("Rate Index Code Options", "MASTER_ONLY"),
                    ],
                )
                self.assertEqual(
                    first_detail_statuses,
                    [
                        ("Match", "Match"),
                        ("Diff", "Not found"),
                        ("Match", "Match"),
                        (None, "Diff"),
                    ],
                )
            finally:
                result_workbook.close()

    def test_combine_result_leaves_one_blank_column_between_files(self):
        workbook = Workbook()
        worksheet = workbook.active
        records = [
            {
                "name": "14_999_999",
                "values": ["MASTER_VALUE"],
                "is_master": True,
            },
            {
                "name": "14_470_991",
                "values": ["VALUE_A"],
                "is_master": False,
            },
            {
                "name": "14_471_991",
                "values": ["VALUE_B"],
                "is_master": False,
            },
        ]

        build_comparison_sheet(worksheet, records)

        self.assertEqual(worksheet["A1"].value, "14_999_999")
        self.assertIsNone(worksheet["B1"].value)
        self.assertEqual(worksheet["C1"].value, "14_470_991")
        self.assertEqual(
            worksheet["D1"].value,
            "Combine Detail to Master",
        )
        self.assertEqual(
            worksheet["E1"].value,
            "Combine Master to Detail",
        )
        self.assertIsNone(worksheet["F1"].value)
        self.assertEqual(worksheet["G1"].value, "14_471_991")
        self.assertEqual(records[1]["result_column"], 3)
        self.assertEqual(records[2]["result_column"], 7)


if __name__ == "__main__":
    unittest.main()
