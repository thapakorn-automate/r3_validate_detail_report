
from collections import OrderedDict
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.combine_detail import extract_not_match_by_file_name


# ============================================================
# CONFIGURATION
# ============================================================

MASTER_FILE_STEM = "14_999_999"

SUMMARY_SHEET_NAME = "Summary"
RESULT_SHEET_NAME = "Combine result"
DETAIL_SHEET_NAME = "Combine detail"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Arial", bold=True, color="000000")
NORMAL_FONT = Font(name="Arial", size=10, color="000000")

MATCH_FILL = PatternFill("solid", fgColor="FFD9EAD3")
MATCH_FONT = Font(name="Arial", size=10, color="FF274E13")

DIFF_FILL = PatternFill("solid", fgColor="FFF4CCCC")
DIFF_FONT = Font(name="Arial", size=10, color="FF990000")

THIN_SIDE = Side(style="thin", color="FFD9E1F2")
CELL_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


# ============================================================
# BASIC HELPERS
# ============================================================

def normalize_value(value):
    return "" if value is None else str(value).strip()


def comparison_key(value):
    return normalize_value(value).casefold()
def display_group_title(file_path):
    import re

    stem = Path(file_path).stem

    # Master title
    if "14_999_999" in stem.casefold():
        return "14_999_999"

    # Example:
    # Detail_Normal_14_470_99H_260922_Diff.xlsx
    # -> 14_470_99H
    match = re.search(
        r"Normal_(.+?)_\d{6,8}(?:_|$)",
        stem,
        re.IGNORECASE,
    )

    if match:
        return match.group(1)

    return stem


def unique_values(values):
    result = []
    seen = set()

    for value in values:
        normalized = normalize_value(value)
        key = comparison_key(normalized)

        if normalized and key not in seen:
            seen.add(key)
            result.append(normalized)

    return result


def is_file_name_marker(value):
    return normalize_value(value).casefold() == "file name"


def get_header_from_filename(filename):
    stem = Path(filename).stem

    match = re.search(
        r"Normal_(.+?)_\d{6}_Diff$",
        stem,
        flags=re.IGNORECASE,
    )

    return match.group(1).strip() if match else stem


# ============================================================
# EXCEL DATA EXTRACTION
# ============================================================

def collect_file_data(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    collected = []

    try:
        for ws in workbook.worksheets:
            for row in ws.iter_rows(
                min_row=1,
                max_row=ws.max_row,
                min_col=1,
                max_col=3,
                values_only=True,
            ):
                if row and is_file_name_marker(row[0]):
                    value = normalize_value(row[2])
                    if value:
                        collected.append(value)
    finally:
        workbook.close()

    return unique_values(collected)


# ============================================================
# FORMATTING HELPERS
# ============================================================

def write_header(ws, row, column, value):
    cell = ws.cell(row=row, column=column, value=value)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.border = CELL_BORDER
    cell.alignment = CENTER
    return cell


def write_header_combined(ws, row, start_column, end_column, value):
    ws.merge_cells(
        start_row=row,
        start_column=start_column,
        end_row=row,
        end_column=end_column,
    )

    cell = ws.cell(row=row, column=start_column, value=value)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.border = CELL_BORDER
    cell.alignment = CENTER

    for column in range(start_column, end_column + 1):
        ws.cell(row=row, column=column).fill = HEADER_FILL
        ws.cell(row=row, column=column).border = CELL_BORDER

    return cell


def write_data_cell(ws, row, column, value):
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = NORMAL_FONT
    cell.border = CELL_BORDER
    cell.alignment = LEFT
    return cell


def write_status_cell(ws, row, column, status):
    cell = ws.cell(row=row, column=column, value=status)
    cell.border = CELL_BORDER
    cell.alignment = CENTER

    if status == "Match":
        cell.fill = MATCH_FILL
        cell.font = MATCH_FONT
    elif status == "Diff":
        cell.fill = DIFF_FILL
        cell.font = DIFF_FONT
    else:
        cell.font = NORMAL_FONT

    return cell


def write_count(ws, row, column, count):
    cell = ws.cell(
        row=row,
        column=column,
        value=f"Count : {count}",
    )
    cell.font = BOLD_FONT
    cell.border = CELL_BORDER
    cell.alignment = CENTER
    return cell


def write_match_diff_counts(
    ws, column, count_row, match_count, diff_count
):
    match_cell = ws.cell(
        row=count_row,
        column=column,
        value=f"Match : {match_count}",
    )
    match_cell.font = BOLD_FONT
    match_cell.border = CELL_BORDER
    match_cell.alignment = CENTER

    diff_cell = ws.cell(
        row=count_row + 1,
        column=column,
        value=f"Diff : {diff_count}",
    )
    diff_cell.font = BOLD_FONT
    diff_cell.border = CELL_BORDER
    diff_cell.alignment = CENTER


def set_internal_hyperlink(cell, sheet_name, target_cell):
    safe_sheet_name = sheet_name.replace("'", "''")
    cell.hyperlink = f"#'{safe_sheet_name}'!{target_cell}"
    cell.style = "Hyperlink"


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    for column_cells in ws.columns:
        if not column_cells:
            continue

        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            (
                len(str(cell.value))
                for cell in column_cells
                if cell.value is not None
            ),
            default=0,
        )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 3, 15), 40
        )


# ============================================================
# SUMMARY SHEET
# ============================================================

def build_summary_sheet(ws, records):
    ws.sheet_view.showGridLines = False

    write_header(ws, 1, 1, "Product")
    write_header(ws, 1, 2, "Count")
    write_header_combined(ws, 1, 3, 4, "Combine Detail to Master")
    write_header_combined(ws, 1, 5, 6, "Combine Master to Detail")

    for column, value in (
        (3, "Match"),
        (4, "Diff"),
        (5, "Match"),
        (6, "Diff"),
    ):
        write_header(ws, 2, column, value)

    for row_index, record in enumerate(records, start=3):
        product_cell = write_data_cell(
            ws, row_index, 1, record["name"]
        )
        write_data_cell(
            ws, row_index, 2, len(record["values"])
        )

        if not record.get("is_master", False):
            keys = (
                "match_detail_to_master",
                "diff_detail_to_master",
                "match_master_to_detail",
                "diff_master_to_detail",
            )

            for column, key in zip(range(3, 7), keys):
                cell = write_data_cell(
                    ws, row_index, column, record.get(key, 0)
                )

                if column in (4, 6) and cell.value:
                    cell.fill = DIFF_FILL
                    cell.font = DIFF_FONT

        result_column = (
            1 if record.get("is_master", False)
            else record["result_column"]
        )

        set_internal_hyperlink(
            product_cell,
            RESULT_SHEET_NAME,
            f"{get_column_letter(result_column)}1",
        )

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:F{max(ws.max_row, 2)}"
    style_sheet(ws)


# ============================================================
# COMBINE RESULT SHEET
# ============================================================

def build_comparison_sheet(ws, records):
    ws.sheet_view.showGridLines = False

    master_record = next(
        record for record in records
        if record.get("is_master", False)
    )

    master_values = master_record["values"]
    master_keys = {comparison_key(value) for value in master_values}

    write_header(ws, 1, 1, master_record["name"])

    for row_index, value in enumerate(master_values, start=2):
        write_data_cell(ws, row_index, 1, value)

    master_diff_append_row = len(master_values) + 2
    current_column = 3

    for record in records:
        if record.get("is_master", False):
            continue

        detail_values = record["values"]
        detail_keys = {comparison_key(value) for value in detail_values}

        value_column = current_column
        detail_to_master_column = current_column + 1
        master_to_detail_column = current_column + 2

        record["result_column"] = value_column

        write_header(ws, 1, value_column, record["name"])
        write_header(
            ws, 1, detail_to_master_column,
            "Combine Detail to Master",
        )
        write_header(
            ws, 1, master_to_detail_column,
            "Combine Master to Detail",
        )

        match_detail_to_master = 0
        diff_detail_to_master = 0
        detail_only_values = []

        for row_index, value in enumerate(detail_values, start=2):
            write_data_cell(ws, row_index, value_column, value)

            if comparison_key(value) in master_keys:
                status = "Match"
                match_detail_to_master += 1
            else:
                status = "Diff"
                diff_detail_to_master += 1
                detail_only_values.append(value)

            write_status_cell(
                ws, row_index, detail_to_master_column, status
            )

        match_master_to_detail = 0
        diff_master_to_detail = 0
        missing_master_values = []

        for value in master_values:
            if comparison_key(value) in detail_keys:
                match_master_to_detail += 1
            else:
                diff_master_to_detail += 1
                missing_master_values.append(value)

        for row_index, value in enumerate(detail_values, start=2):
            status = (
                "Match"
                if comparison_key(value) in master_keys
                else "Diff"
            )
            write_status_cell(
                ws, row_index, master_to_detail_column, status
            )

        for value in detail_only_values:
            cell = write_data_cell(
                ws,
                master_diff_append_row,
                1,
                f"{value} (Diff file : {record['name']})",
            )
            cell.fill = DIFF_FILL
            cell.font = DIFF_FONT
            master_diff_append_row += 1


        record["match_detail_to_master"] = match_detail_to_master
        record["diff_detail_to_master"] = diff_detail_to_master
        record["match_master_to_detail"] = match_master_to_detail
        record["diff_master_to_detail"] = diff_master_to_detail

        count_row = max(
            len(detail_values),
            len(master_values),
        ) + 3

        write_count(ws, count_row, value_column, len(detail_values))

        write_match_diff_counts(
            ws,
            detail_to_master_column,
            count_row,
            match_detail_to_master,
            diff_detail_to_master,
        )
        write_match_diff_counts(
            ws,
            master_to_detail_column,
            count_row,
            match_master_to_detail,
            diff_master_to_detail,
        )

        current_column += 4

    master_count_row = max(
        len(master_values) + 1,
        master_diff_append_row - 1,
    ) + 3

    write_count(ws, master_count_row, 1, len(master_values))

    ws.freeze_panes = "A2"
    style_sheet(ws)



# ============================================================
# SOURCE DATA EXTRACTION FOR COMBINE DETAIL
# ============================================================



def get_unique_file_field_pairs_from_source(file_path):
    """Return unique (File name, Field name) pairs with Not match."""
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    collected = []
    seen_pairs = set()

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    def normalize_header(value):
        value = normalize_value(value).casefold()
        value = re.sub(r"[_\r\n\t]+", " ", value)
        return " ".join(value.split())

    try:
        for sheet in workbook.worksheets:
            current_file_name = ""
            field_col = None
            result_col = None

            for row in sheet.iter_rows(values_only=True):
                column_a = normalize_value(
                    row[0] if len(row) > 0 else ""
                )
                column_b = normalize_value(
                    row[1] if len(row) > 1 else ""
                )

                # พบหัวข้อ File Name ให้เริ่มกลุ่มใหม่
                if is_file_name_marker(column_a):
                    current_file_name = column_b
                    field_col = None
                    result_col = None
                    continue

                headers = {
                    normalize_header(value): index
                    for index, value in enumerate(row)
                    if value is not None
                }

                if (
                    "field name" in headers
                    and "result" in headers
                ):
                    field_col = headers["field name"]
                    result_col = headers["result"]
                    continue

                if field_col is None or result_col is None:
                    continue

                if not current_file_name:
                    continue

                field_name = normalize_value(
                    row[field_col]
                    if field_col < len(row)
                    else ""
                )

                result = normalize_value(
                    row[result_col]
                    if result_col < len(row)
                    else ""
                )

                if (
                    result.casefold() != "not match"
                    or not field_name
                ):
                    continue

                # Unique ตาม File name + Field name
                pair = (
                    current_file_name,
                    field_name,
                )

                pair_key = (
                    comparison_key(current_file_name),
                    comparison_key(field_name),
                )

                if pair_key in seen_pairs:
                    continue

                seen_pairs.add(pair_key)
                collected.append(pair)

    finally:
        workbook.close()

    return collected


def get_unique_field_names_with_not_match(file_path):
    """Return unique Field Name values whose Result is Not match."""
    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    def normalize_header(value):
        value = normalize_value(value).casefold()
        value = re.sub(r"[_\r\n\t]+", " ", value)
        return " ".join(value.split())

    field_names = OrderedDict()

    try:
        for sheet in workbook.worksheets:
            field_col = None
            result_col = None

            for row in sheet.iter_rows(values_only=True):
                headers = {
                    normalize_header(value): index
                    for index, value in enumerate(row)
                    if value is not None
                }

                if "field name" in headers and "result" in headers:
                    field_col = headers["field name"]
                    result_col = headers["result"]
                    continue

                if field_col is None or result_col is None:
                    continue

                field_name = normalize_value(
                    row[field_col] if field_col < len(row) else ""
                )
                result = normalize_value(
                    row[result_col] if result_col < len(row) else ""
                )

                if not field_name or result.casefold() != "not match":
                    continue

                field_names.setdefault(
                    comparison_key(field_name),
                    field_name,
                )
    finally:
        workbook.close()

    return list(field_names.values())


def get_unique_file_names_from_source(file_path):
    """Return unique File names from Column B."""
    file_names = []
    seen_names = set()

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(
                min_col=1,
                max_col=2,
                values_only=True,
            ):
                column_a = normalize_value(row[0])
                column_b = normalize_value(row[1])

                if column_a.casefold() != "file name" or not column_b:
                    continue

                key = column_b.casefold()
                if key not in seen_names:
                    seen_names.add(key)
                    file_names.append(column_b)
    finally:
        workbook.close()

    return file_names


# ============================================================
# COMBINE DETAIL SHEET
# ============================================================


def build_combine_detail_sheet(ws, records):
    """
    Build Combine detail sheet.

    Master:
        - Shows File Name and Not match value only.

    Detail:
        - Shows original Detail values only.
        - Detail values not found in Master -> Not found (gray).
        - Does NOT append missing Master values.
        - Does NOT stamp missing Master File Name / Field Name.
        - Match/Diff counts are calculated from original data only.
    """

    # --------------------------------------------------------
    # HELPERS
    # --------------------------------------------------------
    def normalize(value):
        if value is None:
            return ""

        value = str(value).strip()

        if value.endswith(".0"):
            try:
                value = str(int(float(value)))
            except ValueError:
                pass

        return value

    def comparison_key(value):
        return normalize(value).casefold()

    def collect_grouped_values(file_path):
        grouped = OrderedDict()

        extracted_groups = extract_not_match_by_file_name(
            file_path
        )

        for group in extracted_groups:
            file_name = normalize(group.get("file_name"))

            if not file_name:
                continue

            grouped.setdefault(file_name, [])
            seen = {
                comparison_key(value)
                for value in grouped[file_name]
            }

            for value in group.get("values", []):
                value_text = normalize(value)

                if not value_text:
                    continue

                key = comparison_key(value_text)

                # Unique by File Name + Not match value
                if key not in seen:
                    grouped[file_name].append(value_text)
                    seen.add(key)

        return grouped

    def write_not_found(row, column):
        cell = ws.cell(
            row=row,
            column=column,
            value="Not found",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFD9D9D9",
        )
        cell.font = Font(
            name="Arial",
            size=10,
            color="FF666666",
            bold=True,
        )
        cell.border = CELL_BORDER
        cell.alignment = CENTER
        return cell

    # --------------------------------------------------------
    # RESET ONLY THIS SHEET
    # --------------------------------------------------------
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

    # --------------------------------------------------------
    # ORDER INPUT FILES: MASTER FIRST, THEN DETAILS
    # --------------------------------------------------------
    master_record = None
    detail_records = []

    for record in records:
        file_path = Path(record["path"])

        if MASTER_FILE_STEM.casefold() in file_path.stem.casefold():
            master_record = record
        else:
            detail_records.append(record)

    detail_records.sort(
        key=lambda record: Path(record["path"]).stem.casefold()
    )

    ordered_records = []

    if master_record:
        ordered_records.append(master_record)

    ordered_records.extend(detail_records)

    if not ordered_records:
        cell = ws.cell(
            row=1,
            column=1,
            value="No input files found",
        )
        cell.font = BOLD_FONT
        cell.border = CELL_BORDER
        cell.alignment = CENTER
        return

    # --------------------------------------------------------
    # EXTRACT NOT MATCH VALUES FROM EACH FILE
    # --------------------------------------------------------
    grouped_data = {}

    for record in ordered_records:
        file_path = Path(record["path"])
        grouped_data[str(file_path)] = collect_grouped_values(
            file_path
        )

    master_data = {}

    if master_record:
        master_path = str(Path(master_record["path"]))
        master_data = grouped_data.get(master_path, {})

    # --------------------------------------------------------
    # WRITE HORIZONTAL GROUPS
    # --------------------------------------------------------
    current_column = 1

    for record in ordered_records:
        file_path = Path(record["path"])
        path_key = str(file_path)

        is_master = (
            master_record is not None
            and path_key == str(Path(master_record["path"]))
        )

        detail_data = grouped_data.get(path_key, {})
        group_title = display_group_title(file_path)

        # Master has 2 columns; Detail has 4 columns
        group_width = 2 if is_master else 4

        write_header_combined(
            ws,
            1,
            current_column,
            current_column + group_width - 1,
            group_title,
        )

        if is_master:
            headers = [
                "File Name",
                "Not match value",
            ]
        else:
            headers = [
                "File Name",
                "Not match value",
                "Combine Detail to Master",
                "Combine Master to Detail",
            ]

        for offset, header in enumerate(headers):
            write_header(
                ws,
                2,
                current_column + offset,
                header,
            )

        output_row = 3
        original_count = 0

        detail_match = 0
        detail_diff = 0

        master_match = 0
        master_diff = 0
        master_not_found = 0

        # ----------------------------------------------------
        # MASTER GROUP
        # ----------------------------------------------------
        if is_master:
            for file_name, master_values in master_data.items():
                for value in master_values:
                    write_data_cell(
                        ws,
                        output_row,
                        current_column,
                        file_name,
                    )
                    write_data_cell(
                        ws,
                        output_row,
                        current_column + 1,
                        value,
                    )

                    output_row += 1
                    original_count += 1

        # ----------------------------------------------------
        # DETAIL GROUP
        # ----------------------------------------------------
        else:
            # Preserve Detail File Name order first.
            all_file_names = list(detail_data.keys())

            # Include Master File Names for comparison,
            # but do NOT create rows for missing Master values.
            for file_name in master_data:
                if file_name not in all_file_names:
                    all_file_names.append(file_name)

            for file_name in all_file_names:
                detail_values = detail_data.get(file_name, [])
                master_values = master_data.get(file_name, [])

                detail_keys = {
                    comparison_key(value)
                    for value in detail_values
                }
                master_keys = {
                    comparison_key(value)
                    for value in master_values
                }

                # --------------------------------------------
                # 1. Write original Detail rows only
                # --------------------------------------------

                for value in detail_values:
                    key = comparison_key(value)
                    exists_in_master = key in master_keys

                    if exists_in_master:
                        status_detail_to_master = "Match"
                        status_master_to_detail = "Match"
                        detail_match += 1
                    else:
                        status_detail_to_master = "Diff"
                        status_master_to_detail = "Not found"
                        detail_diff += 1
                        master_not_found += 1


                    write_data_cell(
                        ws,
                        output_row,
                        current_column,
                        file_name,
                    )
                    write_data_cell(
                        ws,
                        output_row,
                        current_column + 1,
                        value,
                    )

                    write_status_cell(
                        ws,
                        output_row,
                        current_column + 2,
                        status_detail_to_master,
                    )

                    if status_master_to_detail == "Not found":
                        write_not_found(
                            output_row,
                            current_column + 3,
                        )
                    else:
                        write_status_cell(
                            ws,
                            output_row,
                            current_column + 3,
                            status_master_to_detail,
                        )

                    output_row += 1
                    original_count += 1

                # --------------------------------------------
                # 2. Calculate Master -> Detail counts only
                #    Do NOT append missing Master values.
                # --------------------------------------------
                for value in master_values:
                    key = comparison_key(value)

                    if key in detail_keys:
                        master_match += 1
                    else:
                        master_diff += 1

        # ----------------------------------------------------
        # COUNTS
        # Only original extracted rows are displayed/countable.
        # Master -> Detail Diff is calculated without stamping
        # missing Master values into the Detail group.
        # ----------------------------------------------------
        count_row = output_row + 1

        count_cell = ws.cell(
            row=count_row,
            column=current_column,
            value="Count",
        )
        count_cell.font = BOLD_FONT
        count_cell.border = CELL_BORDER
        count_cell.alignment = CENTER

        count_value_cell = ws.cell(
            row=count_row,
            column=current_column + 1,
            value=original_count,
        )
        count_value_cell.font = BOLD_FONT
        count_value_cell.border = CELL_BORDER
        count_value_cell.alignment = CENTER

        if not is_master:
            count_values = (
                (
                    current_column + 2,
                    f"Match: {detail_match} | Diff: {detail_diff}",
                ),
                (
                    current_column + 3,
                    f"Match: {master_match} | "
                    # f"Diff: {master_diff} | "
                    f"Not found: {master_not_found}",
                ),
            )

            for column, value in count_values:
                cell = ws.cell(
                    row=count_row,
                    column=column,
                    value=value,
                )
                cell.font = BOLD_FONT
                cell.border = CELL_BORDER
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        # ----------------------------------------------------
        # COLUMN WIDTHS
        # ----------------------------------------------------
        widths = (
            [24, 30]
            if is_master
            else [24, 30, 27, 32]
        )

        for offset, width in enumerate(widths):
            col_letter = get_column_letter(
                current_column + offset
            )
            ws.column_dimensions[col_letter].width = width

        # One blank column between horizontal file groups
        current_column += group_width + 1

    # --------------------------------------------------------
    # FINAL FORMATTING
    # --------------------------------------------------------
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 36


# ============================================================
# MAIN WORKBOOK CREATION
# ============================================================

def create_combined_workbook(master_file, detail_files):
    master_file = Path(master_file)
    detail_files = [Path(path) for path in detail_files]

    if not master_file.exists():
        raise FileNotFoundError(f"Master file not found: {master_file}")

    detail_files = sorted(
        detail_files,
        key=lambda path: path.name.casefold(),
    )

    all_files = [master_file] + [
        path for path in detail_files
        if path.resolve() != master_file.resolve()
    ]

    records = []

    for file_path in all_files:
        is_master = (
            MASTER_FILE_STEM.casefold()
            in file_path.stem.casefold()
        )

        values = collect_file_data(file_path)

        records.append({
            "name": (
                MASTER_FILE_STEM
                if is_master
                else get_header_from_filename(file_path.name)
            ),
            "path": file_path,
            "values": values,
            "is_master": is_master,
            "result_column": 1 if is_master else None,
            "match_detail_to_master": 0,
            "diff_detail_to_master": 0,
            "match_master_to_detail": 0,
            "diff_master_to_detail": 0,
        })

    records.sort(
        key=lambda record: (
            not record["is_master"],
            record["name"].casefold(),
        )
    )

    if not records or not records[0]["is_master"]:
        raise ValueError(
            f"Could not identify Master using '{MASTER_FILE_STEM}'."
        )

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = SUMMARY_SHEET_NAME

    ws_result = wb.create_sheet(RESULT_SHEET_NAME)
    ws_detail = wb.create_sheet(DETAIL_SHEET_NAME)

    build_comparison_sheet(ws_result, records)
    build_combine_detail_sheet(ws_detail, records)
    build_summary_sheet(ws_summary, records)

    metadata = {
        "master_file": str(master_file),
        "detail_files": [
            str(path)
            for path in detail_files
            if path.resolve() != master_file.resolve()
        ],
        "records": records,
        "total_files": len(records),
        "total_master_values": len(records[0]["values"]),
        "master_count": len(records[0]["values"]),
    }

    return wb, metadata