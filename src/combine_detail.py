
from pathlib import Path
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_DIR = BASE_DIR / "input" / "Round14"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_FILE = OUTPUT_DIR / "Combine_detail.xlsx"

SHEET_NAME = "Combine detail"

FILE_NAME_COLUMN = 1       # Column A
VALUE_COLUMN = 3           # Column C
RESULT_COLUMN = 6          # Column F

FILE_NAME_MARKER = "File Name"
RESULT_HEADER = "Result"
NOT_MATCH_VALUE = "Not match"


# ============================================================
# HELPERS
# ============================================================

def normalize(value):
    """Convert a cell value to trimmed text."""
    if value is None:
        return ""
    return str(value).strip()


def normalize_key(value):
    """Case-insensitive comparison key."""
    return normalize(value).casefold()


def is_file_name_marker(value):
    return normalize_key(value) == normalize_key(FILE_NAME_MARKER)


def is_result_header(value):
    return normalize_key(value) == normalize_key(RESULT_HEADER)


def is_not_match(value):
    return normalize_key(value) == normalize_key(NOT_MATCH_VALUE)


# ============================================================
# EXTRACT NOT MATCH VALUES
# ============================================================

def extract_not_match_by_file_name(file_path):
    """
    Scan every worksheet in one report.

    Rules:
    1. Find File Name in Column A.
    2. Read the group name from Column C of that row.
    3. Find Result in Column F in the following rows.
    4. Extract Column C only when Column F is Not match.
    5. Ignore Match and other results.
    6. Remove duplicate values separately for each File Name.
    """

    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Input Excel file not found: {file_path}"
        )

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    grouped_values = OrderedDict()

    try:
        for worksheet in workbook.worksheets:
            current_file_name = ""
            result_header_found = False

            for row in worksheet.iter_rows(
                min_col=1,
                max_col=6,
                values_only=True,
            ):
                column_a = row[FILE_NAME_COLUMN - 1]
                column_c = row[VALUE_COLUMN - 1]
                column_f = row[RESULT_COLUMN - 1]

                # A new File Name section starts here.
                if is_file_name_marker(column_a):
                    current_file_name = normalize(column_c)
                    result_header_found = False

                    if current_file_name:
                        key = normalize_key(current_file_name)

                        if key not in grouped_values:
                            grouped_values[key] = {
                                "file_name": current_file_name,
                                "values": [],
                                "seen": set(),
                            }

                    continue

                if not current_file_name:
                    continue

                # Wait until Column F contains the Result header.
                if not result_header_found:
                    if is_result_header(column_f):
                        result_header_found = True
                    continue

                # Collect only rows marked Not match.
                if not is_not_match(column_f):
                    continue

                value = normalize(column_c)

                if not value:
                    continue

                group_key = normalize_key(current_file_name)
                value_key = normalize_key(value)

                group = grouped_values[group_key]

                # Unique within this File Name only.
                if value_key not in group["seen"]:
                    group["seen"].add(value_key)
                    group["values"].append(value)

    finally:
        workbook.close()

    return [
        {
            "file_name": group["file_name"],
            "values": group["values"],
        }
        for group in grouped_values.values()
    ]


# ============================================================
# CREATE COMBINE DETAIL SHEET
# ============================================================

def write_combine_detail(output_file, report_files):
    """
    Create an Excel workbook containing a Combine detail sheet.

    Column A: File Name
    Column B: Unique values from Column C where Result = Not match
    """

    report_files = [Path(path) for path in report_files]

    if not report_files:
        raise FileNotFoundError(
            "No Excel report files were found in the input folder."
        )

    all_groups = OrderedDict()

    for report_file in sorted(
        report_files,
        key=lambda path: path.name.casefold(),
    ):
        groups = extract_not_match_by_file_name(report_file)

        for group in groups:
            file_name = group["file_name"]
            group_key = normalize_key(file_name)

            if group_key not in all_groups:
                all_groups[group_key] = {
                    "file_name": file_name,
                    "values": [],
                    "seen": set(),
                }

            target = all_groups[group_key]

            # Unique per File Name across all input reports.
            for value in group["values"]:
                value_key = normalize_key(value)

                if value_key not in target["seen"]:
                    target["seen"].add(value_key)
                    target["values"].append(value)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        name="Arial",
        bold=True,
        color="FFFFFF",
    )
    normal_font = Font(
        name="Arial",
        size=10,
        color="000000",
    )
    thin_side = Side(style="thin", color="D9E1F2")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    # Headers
    for column, title in enumerate(
        ["File Name", "Not match value"],
        start=1,
    ):
        cell = worksheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = border

    # Data
    row_index = 2

    for group in all_groups.values():
        for value in group["values"]:
            for column, cell_value in (
                (1, group["file_name"]),
                (2, value),
            ):
                cell = worksheet.cell(
                    row=row_index,
                    column=column,
                    value=cell_value,
                )
                cell.font = normal_font
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                )

            row_index += 1

    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 65

    worksheet.auto_filter.ref = (
        f"A1:B{max(worksheet.max_row, 1)}"
    )

    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook.save(output_file)
    workbook.close()

    total_values = sum(
        len(group["values"])
        for group in all_groups.values()
    )

    print(f"Input reports: {len(report_files)}")
    print(f"File Name groups: {len(all_groups)}")
    print(f"Unique Not match values: {total_values}")
    print(f"Output: {output_file}")

    return output_file


# ============================================================
# MAIN
# ============================================================

def main():
    if not INPUT_DIR.is_dir():
        raise FileNotFoundError(
            f"Input folder not found: {INPUT_DIR}"
        )

    report_files = [
        path
        for path in INPUT_DIR.iterdir()
        if path.is_file()
        and path.suffix.casefold() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
    ]

    write_combine_detail(OUTPUT_FILE, report_files)


if __name__ == "__main__":
    main()